"""Cantidades teóricas por paño — primera losa (ítem 13, losa alivianada con viguetas pretensadas y EPS, h = 25 cm).

Entradas: Plano 03 (viguetas detectadas por p03_viguetas.py), Plano 01 (áreas), Excel PAO 3 (áreas del cómputo).
Salidas:  ../Cantidades_Teoricas_Primera_Losa_por_Pano.csv  y  ../intermedios/resumen_cantidades.json
Todos los supuestos están en el diccionario PAR (con su fuente). Unidades SI; coordenadas del PDF en pt.
Ejecutar:  python3 p01_p03_textos.py && python3 p03_viguetas.py && python3 cantidades_teoricas.py
"""
import csv, json, math, os, runpy, collections
import numpy as np
import openpyxl
from matplotlib.path import Path

BASE = os.path.dirname(os.path.abspath(__file__))
DIR = os.path.abspath(os.path.join(BASE, '..'))
INT = os.path.join(DIR, 'intermedios')
XLSX = '/home/user/TesisSGC/03_Datos_Obra_Edificio_Posgrado/03_Planillas_y_Computos/PAO_03_Planilla_y_Computos_Metricos_Resolucion_de_Contrato.xlsx'

# ------------------------------------------------------------------ parámetros y supuestos
PAR = dict(
    modulo=0.50,            # m entre ejes de viguetas (Plano 02 corte transversal «0.50»; certificado PRETENSA «50 cm»)
    h_total=0.25,           # m (ítem 13 «H=25»; ET; certificado PRETENSA 20 + 5)
    h_eps=0.20,             # m (Plastobloq 20 × 100 × 44, certificado UNICON). Plano 02 acota 0.15 (discrepancia)
    carpeta=0.05,           # m (Plano 02 «e=5 cm»; ET; PRETENSA)
    eps_ancho_sup=0.44,     # m (Plano 02 nervio: 0.50 − 2 × 0.03; UNICON 44 cm)
    eps_ancho_inf=0.40,     # m (Plano 02 nervio: «0.40»)
    eps_largo=1.00,         # m (Plastobloq y ET: 100 cm)
    eps_ancho_min=0.15,     # m: franja de borde con menos ancho útil se llena de hormigón (supuesto)
    vig_masa=17.2,          # kg/m: vigueta pretensada Concretec (ES-CC-1022, Santa Cruz/La Paz/Sucre). Sin dato PRETENSA
    dens_vig=2500.0,        # kg/m³ hormigón pretensado (supuesto usual)
    apoyo=0.05,             # m de empotramiento por extremo (Plano 02, «DETALLE TÍPICO DE APOYO… 5 cm Empotramiento»)
    corr_raster=0.06,       # m: corrección por el grosor del trazo (0,75 pt) en la medición de luces libres
    # nervios transversales (Plano 02: sección 0,10 × h_eps, P1 2φ10 + P2 2φ10 + estribo). Ubicación NO definida en planos
    nerv_luz_min=4.0,       # m: supuesto S1 → un nervio a media luz si la luz libre media ≥ 4,0 m (práctica usual, a confirmar)
    nerv_ancho=0.10,        # m (Plano 02, corte A-A')
    nerv_anclaje=0.30,      # m por extremo de cada barra φ10 (supuesto ≈ 30φ)
    estribo_sep=0.20,       # m (supuesto; el plano dibuja el estribo sin cota)
    estribo_long=0.56,      # m por estribo φ6 (perímetro con recubrimiento 2 cm + ganchos; supuesto)
    # acero
    phi6=0.2220, phi8=0.3946, phi10=0.6165,   # kg/m (masa nominal = π/4 · d² · 7850)
    malla_sep=0.20,         # φ6 c/20 en ambas direcciones (Plano 02; ET)
    malla_anclaje=0.15,     # m por extremo de barra (penetración en viga; supuesto)
    malla_traslape=0.30,    # m por cada barra de 12 m (supuesto ≈ 50φ; factor 12/11,70)
    barra=12.0,
    neg_diam='phi8',        # negativos: 1φ8 por vigueta en cada apoyo continuo (supuesto: el plano dice «de acuerdo a planos»)
    neg_frac=0.25,          # longitud = 0,25·L a cada lado + ancho de viga (supuesto usual)
    # bloque de la ET (100 × 50 × 20): módulo compatible = 0,50 + 0,06 = 0,56 m (misma holgura que 44 → 0,50)
    modulo_et=0.56,
)
PAR.update(json.loads(os.environ.get('CT_PAR', '{}')))   # escenarios de sensibilidad (no sobrescriben el CSV)
SUF = os.environ.get('CT_SUF', '')
A_BLK = PAR.get('A_blk') or (PAR['eps_ancho_sup'] + PAR['eps_ancho_inf']) / 2 * PAR['h_eps']   # m² sección del bloque (0,084)
A_VIG = PAR['vig_masa'] / PAR['dens_vig']                                        # m² sección de la vigueta (0,00688)
ESC = 12.564

# ------------------------------------------------------------------ datos
g = runpy.run_path(os.path.join(BASE, 'p03_viguetas.py'), run_name='lib')
POLY, VIG, is_beam = g['POLY'], g['res'], g['is_beam']
txt = json.load(open(os.path.join(INT, 'p01_p03_textos.json')))
AREA_PLANO = {f'A{k}': v['area_plano'] for k, v in txt['areas'].items()}
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['COMPUTOS MÉTRICOS']
AREA_XL = {}
for r in list(range(1033, 1038)) + list(range(1053, 1089)):
    nm = ws.cell(r, 2).value; val = ws.cell(r, 9).value
    AREA_XL['A' + nm.split()[-1]] = float(val)
assert len(AREA_XL) == 41 and abs(sum(AREA_XL.values()) - 1109.84) < 0.01

# etiquetas de cantidad del Plano 03 asociadas (p01_p03_textos.py) y conteos de componentes
ETIQ = {'A1': '19', 'A2': '19', 'A3': '19', 'A4': '19', 'A5': '19', 'A6': '19', 'A7': '10 (+9 sin etiqueta)',
        'A8': '22', 'A9': '8', 'A10': '8', 'A11': '2', 'A12': '14', 'A13': '4 (+3 acotadas)', 'A14': '14', 'A15': '11',
        'A16': '2', 'A17': '14', 'A18': '— (9 acotadas)', 'A19': '9', 'A20': '14', 'A21': '14', 'A22': '— (3 cotas)',
        'A23': '— (6 acotadas)', 'A24': '14', 'A25': '14', 'A26': '14', 'A27': '— (9 acotadas)', 'A28': '22', 'A29': '8',
        'A30': '8', 'A31': '2', 'A32': '14', 'A33': '4 (+3 acotadas)', 'A34': '20', 'A35': '20', 'A36': '20', 'A37': '4',
        'A38': '20', 'A39': '8', 'A40': '10 (+4 acotadas)', 'A41': '9 (+3 acotadas)'}
CURVOS = {'A1', 'A2', 'A34', 'A35'}          # atravesados por la viga plana curva VP 30x25 (viguetas continuas)
ADOPTA_N = PAR.get('adopta_n') or {'A34': 19, 'A35': 19, 'A36': 19, 'A38': 19}   # etiqueta 20; dibujo 18/18/18/19; geometría 19


def axis_span(k, v):
    """luz entre ejes: tramo de la recta de la vigueta dentro del polígono del paño (líneas dibujadas eje a eje)"""
    pa = Path(POLY[k])
    m = (v['a'] + v['b']) / 2
    ts = np.arange(m - 120, m + 120, 0.1)
    pts = np.c_[ts, np.full_like(ts, v['off'])] if v['o'] == 'H' else np.c_[np.full_like(ts, v['off']), ts]
    ins = pa.contains_points(pts, radius=1e-6) | pa.contains_points(pts, radius=-1e-6)
    i0 = int(np.argmin(abs(ts - m)))
    if not ins[i0]:
        return None
    lo = i0
    while lo > 0 and ins[lo - 1]:
        lo -= 1
    hi = i0
    while hi < len(ts) - 1 and ins[hi + 1]:
        hi += 1
    return (ts[hi] - ts[lo]) / ESC


def scan(v, sign, frac):
    """distancia (m) desde el eje de la vigueta hasta la cara de la viga más próxima, en sentido perpendicular"""
    t = v['a'] + frac * (v['b'] - v['a'])
    for k in range(1, 160):              # hasta 16 pt ≈ 1,27 m
        d = k * 0.1
        px, py = (t, v['off'] + sign * d) if v['o'] == 'H' else (v['off'] + sign * d, t)
        if is_beam(px, py):
            return (d + 0.375) / ESC
    return None


rows_out, panel_info = [], {}
neg_ends = []   # extremos de viguetas para detectar continuidad sobre apoyos
for k in POLY:
    vs = [dict(v) for v in VIG[k]]
    obs = []
    if k == 'A16':   # sin viguetas dibujadas: celda libre 1,28 m (dirección ejes numéricos), luz entre ejes 1,58 m
        for off in (291.6, 297.9):
            vs.append(dict(o='V', off=off, a=305.88, b=321.96, L_libre=1.28 - PAR['corr_raster'], L_tramos=1.28 - PAR['corr_raster'],
                           n_tramos=1, colores=['(etiqueta)'], manual=True, luz_ejes=1.58))
        obs.append('sin viguetas dibujadas; se adopta la etiqueta «2» y la cota 1,60')
    for v in vs:
        v['L_clear'] = v['L_libre'] + PAR['corr_raster']
        v['L_eps'] = v['L_tramos'] + PAR['corr_raster'] if v['n_tramos'] == 1 else v['L_tramos'] + PAR['corr_raster'] * v['n_tramos']
        if 'luz_ejes' not in v:
            v['luz_ejes'] = axis_span(k, v) or v['L_clear'] + 0.275
    if k in ('A9', 'A10', 'A29', 'A30'):   # luz entre ejes según cota roja «1.40» del Plano 03 (eje 5/12 → eje de la VP)
        for v in vs:
            v['luz_ejes'] = 1.40
    if k in CURVOS:   # viguetas continuas a través de la VP curva: luz libre = máxima del paño
        Lmax = max(v['L_clear'] for v in vs); Smax = max(v['luz_ejes'] for v in vs)
        for v in vs:
            v['L_clear'] = Lmax; v['luz_ejes'] = Smax
        obs.append('viga plana curva VP 30x25 atraviesa el paño: viguetas continuas; filas de EPS descuentan la VP')
    n_dib = len(vs)
    vs.sort(key=lambda v: v['off'])
    # ---- filas de EPS (entre viguetas consecutivas y franjas de borde medidas hasta la cara de viga)
    filas = []   # (largo_m, ancho_util_m)
    for i, v in enumerate(vs):
        for sign in (-1, +1):
            j = i + sign
            nb = vs[j] if 0 <= j < len(vs) else None
            gap = abs(nb['off'] - v['off']) / ESC if nb else None
            if nb is not None and gap <= 0.60 + 1e-6:
                if sign == +1:   # fila interior (se cuenta una vez)
                    filas.append(((v['L_eps'] + nb['L_eps']) / 2, PAR['eps_ancho_sup']))
                continue
            e = [scan(v, sign, f) for f in (0.25, 0.5, 0.75)] if not v.get('manual') else [0.30, 0.30, 0.30]
            if None in e or (nb is not None and min(e) >= gap - 0.05):
                continue
            util = np.mean(e) - 0.03
            if util < PAR['eps_ancho_min']:
                continue
            largo = v['L_eps'] if abs(e[0] - e[2]) < 0.08 else v['L_eps'] / 2   # borde paralelo o inclinado
            filas.append((largo, min(PAR['eps_ancho_sup'], util)))
    # ---- ajuste de n adoptado (franja inferior)
    n_adop = ADOPTA_N.get(k, n_dib)
    if n_adop != n_dib:
        Ltip = max(v['L_clear'] for v in vs); Leps = max(v['L_eps'] for v in vs); S = max(v['luz_ejes'] for v in vs)
        for _ in range(n_adop - n_dib):
            vs.append(dict(o=vs[0]['o'], off=None, L_clear=Ltip, L_eps=Leps, luz_ejes=S, n_tramos=1, adoptada=True))
        # filas: el ancho entre vigas no cambia → (n+1) filas de módulo completo (las de borde se reducen)
        filas = [(f[0], f[1]) for f in filas]
        while len(filas) < n_adop + 1:
            filas.append((Leps, PAR['eps_ancho_sup']))
        obs.append(f'etiqueta 20, dibujadas {n_dib}; se adopta {n_adop} (10,00 m entre ejes, módulo 0,50 m, igual a la franja superior)')
    n = len(vs)
    L_clear = [v['L_clear'] for v in vs]
    luz = [v['luz_ejes'] for v in vs]
    long_vig = [L + 2 * PAR['apoyo'] for L in L_clear]
    # ---- nervios transversales (S1)
    Lm = float(np.mean(L_clear))
    n_nerv = 1 if Lm >= PAR['nerv_luz_min'] else 0
    Wperp = AREA_XL[k] / Lm if Lm > 0 else 0
    # ---- EPS (Plastobloq 20 × 100 × 44)
    L_filas = [f[0] for f in filas]
    eps_net_S0 = sum(L_filas) / PAR['eps_largo']
    eps_net_S1 = sum(max(0.0, L - n_nerv * PAR['nerv_ancho']) for L in L_filas) / PAR['eps_largo']
    eps_pra_S0 = sum(math.ceil(L / PAR['eps_largo'] - 1e-9) for L in L_filas)
    eps_pra_S1 = sum(math.ceil(max(0.0, L - n_nerv * PAR['nerv_ancho']) / PAR['eps_largo'] - 1e-9) for L in L_filas)
    filas_et = max(1, round(len(filas) * PAR['modulo'] / PAR['modulo_et'])) if filas else 0
    eps_et_S1 = filas_et * (np.mean([max(0.0, L - n_nerv * PAR['nerv_ancho']) for L in L_filas]) if filas else 0) / PAR['eps_largo']
    # ---- hormigón
    V_eps_S0 = sum(L * A_BLK * min(1.0, w / PAR['eps_ancho_sup']) for L, w in filas)
    V_eps_S1 = sum(max(0.0, L - n_nerv * PAR['nerv_ancho']) * A_BLK * min(1.0, w / PAR['eps_ancho_sup']) for L, w in filas)
    V_vig = sum(L_clear) * A_VIG
    V_hor_S0 = AREA_XL[k] * PAR['h_total'] - V_eps_S0 - V_vig
    V_hor_S1 = AREA_XL[k] * PAR['h_total'] - V_eps_S1 - V_vig
    V_carpeta = AREA_XL[k] * PAR['carpeta']
    # ---- acero: malla φ6 c/20 en dos direcciones (+ anclajes + traslapes)
    pa = POLY[k]
    per = sum(math.dist(pa[i], pa[(i + 1) % len(pa)]) for i in range(len(pa))) / ESC
    L_malla = (2 * AREA_XL[k] / PAR['malla_sep'] + (per / 2 / PAR['malla_sep']) * 2 * PAR['malla_anclaje']) * PAR['barra'] / (PAR['barra'] - PAR['malla_traslape'])
    kg_malla = L_malla * PAR['phi6']
    # nervios: 4φ10 + estribos φ6
    kg_nerv = n_nerv * (4 * (Wperp + 2 * PAR['nerv_anclaje']) * PAR['phi10'] + (math.ceil(Wperp / PAR['estribo_sep']) + 1) * PAR['estribo_long'] * PAR['phi6'])
    panel_info[k] = dict(n=n, n_dib=n_dib, filas=len(filas), Lm=Lm, n_nerv=n_nerv, Wperp=Wperp, kg_malla=kg_malla, kg_nerv=kg_nerv,
                         eps_net_S0=eps_net_S0, eps_net_S1=eps_net_S1, eps_pra_S0=eps_pra_S0, eps_pra_S1=eps_pra_S1, eps_et_S1=eps_et_S1,
                         V_hor_S0=V_hor_S0, V_hor_S1=V_hor_S1, V_carpeta=V_carpeta, V_eps_S1=V_eps_S1, V_vig=V_vig,
                         m_vig=sum(long_vig), luz_media=float(np.mean(luz)), Lclear_media=Lm, long_media=float(np.mean(long_vig)),
                         dir=('ejes literales (A→M)' if vs[0]['o'] == 'H' else 'ejes numéricos (1→16)'), obs=obs,
                         L_lista=sorted(round(x, 2) for x in long_vig), luz_lista=sorted(round(x, 2) for x in luz))
    for v in vs:
        if v.get('off') is not None and not v.get('manual'):
            neg_ends.append((k, v['o'], v['off'], v['a'], v['b'], v['L_clear']))

# ------------------------------------------------------------------ negativos: continuidad sobre apoyos
neg_kg = collections.defaultdict(float); n_pares = 0; pares = []
phi_neg = PAR[PAR['neg_diam']]
for i, (k1, o1, off1, a1, b1, L1) in enumerate(neg_ends):
    for j in range(i + 1, len(neg_ends)):
        k2, o2, off2, a2, b2, L2 = neg_ends[j]
        if k1 == k2 or o1 != o2 or abs(off1 - off2) > 0.8:
            continue
        gap = min(abs(a2 - b1), abs(a1 - b2)) / ESC
        if gap <= 0.45:
            Lb = PAR['neg_frac'] * (L1 + L2) + gap
            neg_kg[k1] += Lb * phi_neg / 2; neg_kg[k2] += Lb * phi_neg / 2; n_pares += 1
            pares.append((k1, k2))
for k, adop in ADOPTA_N.items():   # viguetas añadidas: negativos proporcionales
    pi = panel_info[k]
    neg_kg[k] *= pi['n'] / pi['n_dib']

# ------------------------------------------------------------------ CSV
csv_path = os.path.join(DIR, 'Cantidades_Teoricas_Primera_Losa_por_Pano.csv') if not SUF else os.devnull
with open(csv_path, 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh)
    w.writerow(['pano', 'area_m2_excel', 'area_m2_plano', 'n_viguetas', 'luz_vigueta_m', 'long_vigueta_m', 'm_vigueta',
                'pza_eps_plastobloq', 'pza_eps_et', 'm3_hormigon', 'kg_acero', 'observaciones'])
    for k in POLY:
        pi = panel_info[k]
        kg = pi['kg_malla'] + neg_kg[k] + pi['kg_nerv']
        o = [f"etiqueta Plano 03: {ETIQ[k]}", f"dirección viguetas: {pi['dir']}", f"filas EPS: {pi['filas']}",
             f"nervio transversal (S1): {pi['n_nerv']}"]
        if max(pi['L_lista']) - min(pi['L_lista']) > 0.15:
            o.append('longitudes de vigueta (m): ' + '; '.join(f'{x:.2f}' for x in pi['L_lista']))
        o.append(f"luz libre media {pi['Lclear_media']:.2f} m; long. = luz libre + 2 × 0,05 m")
        o += pi['obs']
        if abs(AREA_XL[k] - AREA_PLANO[k]) > 0.005:
            o.append(f'área Plano 01 ≠ Excel ({AREA_PLANO[k]} vs {AREA_XL[k]})')
        w.writerow([k, f'{AREA_XL[k]:.2f}', f'{AREA_PLANO[k]:.2f}', pi['n'], f"{pi['luz_media']:.2f}", f"{pi['long_media']:.2f}",
                    f"{pi['m_vig']:.2f}", f"{pi['eps_net_S1']:.1f}", f"{pi['eps_et_S1']:.1f}", f"{pi['V_hor_S1']:.3f}", f'{kg:.1f}',
                    ' | '.join(o)])

# ------------------------------------------------------------------ resumen
A = sum(AREA_XL.values())
tot = lambda key: sum(panel_info[k][key] for k in POLY)
kg_neg = sum(neg_kg.values())
R = dict(area=A, n_vig=tot('n'), m_vig=tot('m_vig'), eps_net_S0=tot('eps_net_S0'), eps_net_S1=tot('eps_net_S1'),
         eps_pra_S0=tot('eps_pra_S0'), eps_pra_S1=tot('eps_pra_S1'), eps_et_S1=tot('eps_et_S1'),
         V_hor_S0=tot('V_hor_S0'), V_hor_S1=tot('V_hor_S1'), V_carpeta=tot('V_carpeta'), V_vig=tot('V_vig'),
         kg_malla=tot('kg_malla'), kg_neg=kg_neg, kg_nerv=tot('kg_nerv'), n_nerv=tot('n_nerv'), n_pares_neg=n_pares,
         A_BLK=A_BLK, A_VIG=A_VIG, PAR=PAR)
R['kg_S1'] = R['kg_malla'] + R['kg_neg'] + R['kg_nerv']
R['kg_S0'] = R['kg_malla'] + R['kg_neg']
for key in ['n_vig', 'm_vig', 'eps_net_S0', 'eps_net_S1', 'eps_pra_S0', 'eps_pra_S1', 'eps_et_S1', 'V_hor_S0', 'V_hor_S1', 'kg_S0', 'kg_S1', 'kg_malla', 'kg_neg', 'kg_nerv']:
    R[key + '_m2'] = R[key] / A
EPS_REAL = 2350
for key in ['eps_net_S0', 'eps_net_S1', 'eps_pra_S0', 'eps_pra_S1']:
    R['IDM_' + key] = (EPS_REAL - R[key]) / R[key] * 100
for key in ['kg_S0', 'kg_S1']:
    R['dif_4171_' + key] = (4171.18 - R[key]) / R[key] * 100
    R['dif_12513_' + key] = (12513.54 - R[key]) / R[key] * 100
R['segunda_tercera_x2'] = {k: 2 * R[k] for k in ['n_vig', 'm_vig', 'eps_net_S1', 'eps_pra_S1', 'V_hor_S1', 'kg_S1', 'kg_S0', 'area']}
R['saldo_2234_60'] = {k: 2234.60 * R[k + '_m2'] for k in ['n_vig', 'm_vig', 'eps_net_S1', 'V_hor_S1', 'kg_S1', 'kg_S0']}
R['por_pano'] = {k: {kk: (vv if not isinstance(vv, (np.floating,)) else float(vv)) for kk, vv in panel_info[k].items()} for k in POLY}
R['neg_kg_pano'] = dict(neg_kg)
json.dump(R, open(os.path.join(INT, f'resumen_cantidades{SUF}.json'), 'w'), indent=1, ensure_ascii=False, default=float)
print(f"Área {A:.2f} m² | viguetas {R['n_vig']} pza, {R['m_vig']:.1f} m ({R['m_vig_m2']:.3f} m/m²)")
print(f"EPS neto S0 {R['eps_net_S0']:.1f} | S1 {R['eps_net_S1']:.1f} ({R['eps_net_S1_m2']:.3f}/m²) | práctico S0 {R['eps_pra_S0']} S1 {R['eps_pra_S1']} | ET S1 {R['eps_et_S1']:.1f}")
print(f"IDM vs 2350: neto S0 {R['IDM_eps_net_S0']:.1f} %, neto S1 {R['IDM_eps_net_S1']:.1f} %, práctico S0 {R['IDM_eps_pra_S0']:.1f} %, práctico S1 {R['IDM_eps_pra_S1']:.1f} %")
print(f"Hormigón S0 {R['V_hor_S0']:.2f} m³ ({R['V_hor_S0_m2']:.4f}/m²) | S1 {R['V_hor_S1']:.2f} ({R['V_hor_S1_m2']:.4f}/m²) | carpeta {R['V_carpeta']:.2f}")
print(f"Acero malla {R['kg_malla']:.1f} | negativos {R['kg_neg']:.1f} ({n_pares} pares) | nervios {R['kg_nerv']:.1f} ({R['n_nerv']} nervios)")
print(f"Acero S0 {R['kg_S0']:.1f} kg ({R['kg_S0_m2']:.2f}/m²) | S1 {R['kg_S1']:.1f} kg ({R['kg_S1_m2']:.2f}/m²)")
print(f"4171,18 vs S1: {R['dif_4171_kg_S1']:.1f} % ; vs S0: {R['dif_4171_kg_S0']:.1f} % | 12513,54 vs S1: {R['dif_12513_kg_S1']:.1f} %")
