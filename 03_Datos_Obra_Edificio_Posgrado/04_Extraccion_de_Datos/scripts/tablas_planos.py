"""Tablas de trazabilidad (Tarea B): áreas del Plano 01 con coordenadas vs Excel, y etiquetas/cotas del Plano 03.
Requiere p01_p03_textos.py (JSON) y p03_viguetas.py. Salidas en ../intermedios/"""
import csv, json, os, math
import openpyxl
BASE = os.path.dirname(os.path.abspath(__file__)); INT = os.path.join(BASE, '..', 'intermedios')
T = json.load(open(os.path.join(INT, 'p01_p03_textos.json')))
ws = openpyxl.load_workbook('/home/user/TesisSGC/03_Datos_Obra_Edificio_Posgrado/03_Planillas_y_Computos/PAO_03_Planilla_y_Computos_Metricos_Resolucion_de_Contrato.xlsx', data_only=True)['COMPUTOS MÉTRICOS']
XL = {}
for r in list(range(1033, 1038)) + list(range(1053, 1089)):
    XL[int(ws.cell(r, 2).value.split()[-1])] = (r, float(ws.cell(r, 9).value))
with open(os.path.join(INT, 'plano01_areas_vs_excel.csv'), 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh)
    w.writerow(['pano', 'area_plano01_m2', 'x_pdf01_pt', 'y_pdf01_pt', 'x_plano03_pt', 'y_plano03_pt', 'fila_excel', 'area_excel_m2', 'diferencia_m2'])
    for k in sorted(T['areas'], key=int):
        a = T['areas'][k]; r, v = XL[int(k)]
        w.writerow([f'A{k}', a['area_plano'], a['x01'], a['y01'], a['x03'], a['y03'], r, v, round(a['area_plano'] - v, 2)])
VIG = json.load(open(os.path.join(INT, 'p03_viguetas_por_pano.json')))['viguetas']
with open(os.path.join(INT, 'plano03_etiquetas_y_cotas.csv'), 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh)
    w.writerow(['tipo', 'texto', 'x_pt', 'y_pt', 'pano_asociado', 'distancia_pt'])
    for c in T['counts']:
        best = min(T['areas'], key=lambda k: math.hypot(T['areas'][k]['x03'] - c['x'], T['areas'][k]['y03'] - c['y']))
        d = math.hypot(T['areas'][best]['x03'] - c['x'], T['areas'][best]['y03'] - c['y'])
        w.writerow(['cantidad_viguetas', c['n'], c['x'], c['y'], f'A{best}', round(d, 1)])
    for c in T['reds']:
        w.writerow(['cota_roja_longitud', c['t'], c['x'], c['y'], '', ''])
    for c in T['cotas']:
        w.writerow(['cota_negra_entre_ejes', c['t'], c['x'], c['y'], '', ''])
print('ok')
