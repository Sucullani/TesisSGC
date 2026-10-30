# -*- coding: utf-8 -*-
"""Plantilla APA 7 para pandoc y posproceso del .docx con python-docx.

Formato (Guía de Normas APA 7.ª ed. del repositorio, pp. 17-33):
papel carta; márgenes de 2,54 cm; Times New Roman 12; interlineado doble sin
espacio entre párrafos; alineado a la izquierda; sangría de primera línea de
1,27 cm; número de página arriba a la derecha; títulos: nivel 1 centrado en
negrita, nivel 2 a la izquierda en negrita, nivel 3 a la izquierda en negrita
y cursiva; tablas con bordes solo horizontales; «Tabla N»/«Figura N» en
negrita y título en cursiva en la línea siguiente; «Nota.» en cursiva;
referencias con sangría francesa de 1,27 cm.
"""

import copy
import math
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

FUENTE = 'Times New Roman'
FUENTE_CODIGO = 'Courier New'
ANCHO_UTIL_TW = 9360          # 16,51 cm en twips
ALTO_UTIL_TW = 12960          # 22,86 cm
CM_TW = 1440 / 2.54
ROJO_PENDIENTE = RGBColor(0xB6, 0x32, 0x1C)   # BrickRed, como en el PDF de LaTeX


# --------------------------------------------------------------------------
#  Utilidades XML
# --------------------------------------------------------------------------

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace('_', ':', 1)), str(v))
    return e


def hijo(padre, tag, crear=True, antes_de=None):
    h = padre.find(qn(tag))
    if h is None and crear:
        h = OxmlElement(tag)
        if antes_de is not None:
            ref = None
            for t in antes_de:
                ref = padre.find(qn(t))
                if ref is not None:
                    break
            if ref is not None:
                ref.addprevious(h)
            else:
                padre.append(h)
        else:
            padre.append(h)
    return h


def es_tabla_ecuacion(tbl):
    tp = tbl.find(qn('w:tblPr'))
    cap = tp.find(qn('w:tblCaption')) if tp is not None else None
    return cap is not None and cap.get(qn('w:val')) == 'Ecuación'


def texto_p(p):
    return ''.join(t.text or '' for t in p.iter(qn('w:t')))


def pPr(p):
    pp = p.find(qn('w:pPr'))
    if pp is None:
        pp = OxmlElement('w:pPr')
        p.insert(0, pp)
    return pp


# Orden de los hijos de w:pPr según el esquema (para no producir XML inválido)
ORDEN_PPR = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
             'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
             'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
             'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
             'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
             'sectPr', 'pPrChange']


def poner_ppr(p, tag, **attrs):
    """Crea o reemplaza un hijo de w:pPr respetando el orden del esquema."""
    pp = pPr(p)
    viejo = pp.find(qn('w:' + tag))
    if viejo is not None:
        pp.remove(viejo)
    nuevo = el('w:' + tag, **attrs)
    idx = ORDEN_PPR.index(tag)
    for h in pp:
        nombre = h.tag.split('}')[1]
        if nombre in ORDEN_PPR and ORDEN_PPR.index(nombre) > idx:
            h.addprevious(nuevo)
            return nuevo
    pp.append(nuevo)
    return nuevo


def estilo_de(p):
    pp = p.find(qn('w:pPr'))
    if pp is None:
        return None
    ps = pp.find(qn('w:pStyle'))
    return ps.get(qn('w:val')) if ps is not None else None


def quitar_en_texto(p, patron, reemplazo=''):
    """Elimina (o reemplaza) un patrón dentro de los w:t de un párrafo."""
    rx = re.compile(patron)
    for t in p.iter(qn('w:t')):
        if t.text and rx.search(t.text):
            t.text = rx.sub(reemplazo, t.text)
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


# --------------------------------------------------------------------------
#  Plantilla de referencia
# --------------------------------------------------------------------------

def _rfonts(rpr, nombre):
    rf = hijo(rpr, 'w:rFonts')
    for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn('w:' + a), nombre)
    for a in ('asciiTheme', 'hAnsiTheme', 'cstheme', 'eastAsiaTheme'):
        if rf.get(qn('w:' + a)) is not None:
            del rf.attrib[qn('w:' + a)]


def _estilo(doc, nombre, tipo=WD_STYLE_TYPE.PARAGRAPH, base=None):
    st = None
    for s in doc.styles:
        if s.name is not None and s.name.lower() == nombre.lower():
            st = s
            break
    if st is None:
        st = doc.styles.add_style(nombre, tipo)
    if base is not None:
        st.base_style = _estilo(doc, base)
    return st


def _parrafo(st, alineacion=None, interlineado=2.0, antes=0, despues=0, sangria_1=None, izquierda=None,
             colgante=None, mantener_sig=None, salto_antes=None, nivel=None, derecha=None):
    pf = st.paragraph_format
    if alineacion is not None:
        pf.alignment = alineacion
    if interlineado is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE if interlineado not in (1.0, 2.0) else (
            WD_LINE_SPACING.SINGLE if interlineado == 1.0 else WD_LINE_SPACING.DOUBLE)
        if interlineado not in (1.0, 2.0):
            pf.line_spacing = interlineado
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    if colgante is not None:
        pf.left_indent = Cm(izquierda or colgante)
        pf.first_line_indent = Cm(-colgante)
    else:
        if izquierda is not None:
            pf.left_indent = Cm(izquierda)
        if sangria_1 is not None:
            pf.first_line_indent = Cm(sangria_1)
    if derecha is not None:
        pf.right_indent = Cm(derecha)
    if mantener_sig is not None:
        pf.keep_with_next = mantener_sig
    if salto_antes is not None:
        pf.page_break_before = salto_antes
    if nivel is not None:
        ppr = st.element.get_or_add_pPr()
        ol = ppr.find(qn('w:outlineLvl'))
        if ol is None:
            ol = OxmlElement('w:outlineLvl')
            ppr.append(ol)
        ol.set(qn('w:val'), str(nivel))
    pf.widow_control = True


def _letra(st, tam=12, negrita=None, cursiva=None, mayus=None, color=RGBColor(0, 0, 0), fuente=FUENTE):
    f = st.font
    f.name = fuente
    _rfonts(st.element.get_or_add_rPr(), fuente)
    f.size = Pt(tam)
    if negrita is not None:
        f.bold = negrita
    if cursiva is not None:
        f.italic = cursiva
    if mayus is not None:
        f.all_caps = mayus
    if color is not None:
        f.color.rgb = color


def _campo_pagina(parrafo):
    for tipo, texto in (('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)):
        r = parrafo.add_run()
        if tipo:
            r._r.append(el('w:fldChar', w_fldCharType=tipo))
        elif texto == ' PAGE ':
            it = el('w:instrText')
            it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            it.text = texto
            r._r.append(it)
        else:
            r.text = texto


def crear_plantilla(destino, dir_trabajo):
    """Crea plantilla_apa.docx a partir de la plantilla por defecto de pandoc."""
    base = Path(dir_trabajo) / 'reference_pandoc.docx'
    with open(base, 'wb') as fh:
        subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'], stdout=fh, check=True)
    doc = Document(str(base))

    # ---- página
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.orientation = WD_ORIENT.PORTRAIT
    for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, lado, Cm(2.54))
    sec.header_distance = Cm(1.27)
    sec.footer_distance = Cm(1.27)
    sec.gutter = Cm(0)
    sec.header.is_linked_to_previous = False
    ph = sec.header.paragraphs[0]
    ph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ph.style = _estilo(doc, 'Normal')
    ph.paragraph_format.first_line_indent = Cm(0)
    ph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _campo_pagina(ph)

    # ---- valores por defecto del documento
    estilos = doc.styles.element
    rpr_def = estilos.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    _rfonts(rpr_def, FUENTE)
    hijo(rpr_def, 'w:sz').set(qn('w:val'), '24')
    hijo(rpr_def, 'w:szCs').set(qn('w:val'), '24')
    lang = hijo(rpr_def, 'w:lang')
    lang.set(qn('w:val'), 'es-BO')
    lang.set(qn('w:eastAsia'), 'es-BO')
    ppr_def = estilos.find(qn('w:docDefaults')).find(qn('w:pPrDefault'))
    if ppr_def is not None:
        sp = ppr_def.find(qn('w:pPr')).find(qn('w:spacing'))
        if sp is not None:
            sp.set(qn('w:after'), '0')

    I = WD_ALIGN_PARAGRAPH
    # ---- párrafos de texto
    st = _estilo(doc, 'Normal')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT)
    for nombre in ('Body Text', 'First Paragraph'):
        st = _estilo(doc, nombre)
        _letra(st)
        _parrafo(st, alineacion=I.LEFT, sangria_1=1.27)
    st = _estilo(doc, 'Compact')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, sangria_1=0)
    for nombre in ('Title', 'Subtitle', 'Author', 'Date', 'Abstract Title', 'Abstract'):
        st = _estilo(doc, nombre)
        _letra(st, negrita=(nombre == 'Title'), cursiva=False)
        _parrafo(st, alineacion=I.CENTER, sangria_1=0)

    # ---- títulos (niveles APA)
    for n in range(1, 10):
        st = _estilo(doc, 'Heading %d' % n)
        st.base_style = _estilo(doc, 'Normal')
        if n == 1:
            _letra(st, negrita=True, cursiva=False, mayus=True)
            _parrafo(st, alineacion=I.CENTER, sangria_1=0, mantener_sig=True, salto_antes=True, nivel=0)
        elif n == 2:
            _letra(st, negrita=True, cursiva=False, mayus=False)
            _parrafo(st, alineacion=I.LEFT, sangria_1=0, mantener_sig=True, salto_antes=False, nivel=1)
        elif n == 3:
            _letra(st, negrita=True, cursiva=True, mayus=False)
            _parrafo(st, alineacion=I.LEFT, sangria_1=0, mantener_sig=True, salto_antes=False, nivel=2)
        else:
            _letra(st, negrita=True, cursiva=(n >= 5), mayus=False)
            _parrafo(st, alineacion=I.LEFT, izquierda=0, sangria_1=1.27, mantener_sig=True, salto_antes=False,
                     nivel=n - 1)
    st = _estilo(doc, 'TOC Heading')
    st.base_style = _estilo(doc, 'Normal')
    _letra(st, negrita=True, cursiva=False, mayus=True)
    _parrafo(st, alineacion=I.CENTER, sangria_1=0, mantener_sig=True, salto_antes=True, despues=12)
    ol = st.element.get_or_add_pPr().find(qn('w:outlineLvl'))
    if ol is not None:
        ol.getparent().remove(ol)

    # ---- bibliografía, citas largas, notas al pie
    st = _estilo(doc, 'Bibliography')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, colgante=1.27)
    st = _estilo(doc, 'Block Text')
    _letra(st, cursiva=False)
    _parrafo(st, alineacion=I.LEFT, izquierda=1.27, sangria_1=0)
    st.element.get_or_add_pPr()
    pbdr = st.element.pPr.find(qn('w:pBdr'))
    if pbdr is not None:
        st.element.pPr.remove(pbdr)
    st = _estilo(doc, 'Footnote Text')
    _letra(st, tam=10)
    _parrafo(st, alineacion=I.LEFT, interlineado=1.0, sangria_1=0)
    for nombre in ('Definition Term', 'Definition'):
        st = _estilo(doc, nombre)
        _letra(st, negrita=(nombre == 'Definition Term'))
        _parrafo(st, alineacion=I.LEFT, sangria_1=0, izquierda=(0 if nombre == 'Definition Term' else 1.27))

    # ---- tablas y figuras
    for nombre in ('Caption', 'Table Caption', 'Image Caption'):
        st = _estilo(doc, nombre)
        _letra(st, negrita=False, cursiva=False)
        _parrafo(st, alineacion=I.LEFT, sangria_1=0, mantener_sig=True, antes=12)
    for nombre in ('Figure', 'Captioned Figure'):
        st = _estilo(doc, nombre)
        _letra(st)
        _parrafo(st, alineacion=I.CENTER, interlineado=1.0, sangria_1=0, mantener_sig=True)
    st = _estilo(doc, 'Nota APA', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, sangria_1=0, despues=12)
    st = _estilo(doc, 'Figura APA', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.CENTER, interlineado=1.0, sangria_1=0, mantener_sig=True, antes=3, despues=3)
    st = _estilo(doc, 'Ecuacion APA', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.CENTER, interlineado=1.0, sangria_1=0, antes=6, despues=6)
    st = _estilo(doc, 'Celda APA', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, interlineado=1.0, sangria_1=0, antes=1, despues=1)

    # ---- carátula, listas con etiqueta, recuadros
    st = _estilo(doc, 'Caratula', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.CENTER, sangria_1=0)
    st = _estilo(doc, 'Lista etiquetada', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, izquierda=2.54, colgante=1.27)
    st.paragraph_format.tab_stops.add_tab_stop(Cm(2.54))
    st = _estilo(doc, 'Recuadro', base='Normal')
    _letra(st)
    _parrafo(st, alineacion=I.LEFT, interlineado=1.15, sangria_1=0, izquierda=0.4, derecha=0.4, antes=4, despues=4)
    ppr = st.element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    for lado in ('top', 'left', 'bottom', 'right'):
        bdr.append(el('w:' + lado, w_val='single', w_sz='6', w_space='6', w_color='808080'))
    ppr.insert(0, bdr)
    ppr.insert(1, el('w:shd', w_val='clear', w_color='auto', w_fill='F2F2F2'))

    # ---- código fuente
    st = _estilo(doc, 'Source Code', base='Normal')
    _letra(st, tam=9, fuente=FUENTE_CODIGO)
    _parrafo(st, alineacion=I.LEFT, interlineado=1.0, sangria_1=0, antes=6, despues=6)
    st = _estilo(doc, 'Verbatim Char', tipo=WD_STYLE_TYPE.CHARACTER)
    st.font.name = FUENTE_CODIGO
    _rfonts(st.element.get_or_add_rPr(), FUENTE_CODIGO)
    st.font.size = Pt(10)

    # ---- marcas de borrador
    st = _estilo(doc, 'Pendiente', tipo=WD_STYLE_TYPE.CHARACTER)
    st.font.color.rgb = ROJO_PENDIENTE

    # ---- índices (general, de tablas y de figuras)
    for n, izq in ((1, 0.0), (2, 0.6), (3, 1.4)):
        st = _estilo(doc, 'toc %d' % n)
        st.base_style = _estilo(doc, 'Normal')
        _letra(st, negrita=(n == 1), cursiva=False)
        _parrafo(st, alineacion=I.LEFT, interlineado=1.0, izquierda=izq + 0.0, sangria_1=0, derecha=1.0,
                 antes=(6 if n == 1 else 2), despues=2)
        st.paragraph_format.tab_stops.add_tab_stop(Cm(16.51), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    st = _estilo(doc, 'table of figures')
    st.base_style = _estilo(doc, 'Normal')
    _letra(st, negrita=False, cursiva=False)
    _parrafo(st, alineacion=I.LEFT, interlineado=1.0, izquierda=1.9, colgante=1.9, derecha=1.0, antes=2, despues=4)
    st.paragraph_format.tab_stops.add_tab_stop(Cm(16.51), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # ---- hipervínculos (APA 7: el formato por defecto del procesador de texto es válido)
    st = _estilo(doc, 'Hyperlink', tipo=WD_STYLE_TYPE.CHARACTER)
    st.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)

    # ---- estilo de tabla de pandoc: sin bordes, márgenes de celda estándar
    st = _estilo(doc, 'Table', tipo=WD_STYLE_TYPE.TABLE)
    tblpr = st.element.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = OxmlElement('w:tblPr')
        st.element.append(tblpr)
    for viejo in tblpr.findall(qn('w:tblBorders')):
        tblpr.remove(viejo)
    for tbl_style_pr in st.element.findall(qn('w:tblStylePr')):
        st.element.remove(tbl_style_pr)
    doc.save(str(destino))
    return destino


# --------------------------------------------------------------------------
#  Posproceso
# --------------------------------------------------------------------------

class Posproceso:
    def __init__(self, ruta_docx, datos, avisos, tabla_pt=12, formulario_pt=10):
        self.ruta = Path(ruta_docx)
        self.doc = Document(str(ruta_docx))
        self.body = self.doc.element.body
        self.datos = datos
        self.avisos = avisos
        self.tabla_pt = tabla_pt
        self.formulario_pt = formulario_pt
        ids = [int(b.get(qn('w:id'))) for b in self.body.iter(qn('w:bookmarkStart')) if b.get(qn('w:id'), '').isdigit()]
        self.siguiente_bm = max(ids + [0]) + 1000
        self.tablas_ecuacion = set()

    # ------------------------------------------------------------------
    def parrafos(self):
        return list(self.body.iter(qn('w:p')))

    def bloques(self):
        return [e for e in self.body if e.tag in (qn('w:p'), qn('w:tbl'))]

    def marcador(self, patron):
        rx = re.compile(patron)
        res = []
        for p in self.body.iter(qn('w:p')):
            t = texto_p(p)
            m = rx.search(t)
            if m:
                res.append((p, m))
        return res

    def eliminar(self, e):
        padre = e.getparent()
        if padre is not None:
            padre.remove(e)

    def marcador_bm(self, p, nombre):
        bid = str(self.siguiente_bm)
        self.siguiente_bm += 1
        ini = el('w:bookmarkStart', w_id=bid, w_name=nombre)
        fin = el('w:bookmarkEnd', w_id=bid)
        pp = p.find(qn('w:pPr'))
        if pp is not None:
            pp.addnext(ini)
        else:
            p.insert(0, ini)
        p.append(fin)

    # ------------------------------------------------------------------
    def ejecutar(self):
        self.marcadores_bm()
        self.titulos()
        self.tablas()
        self.ecuaciones()
        self.sin_sangria()
        self.tabuladores()
        self.indices()
        self.imagenes()
        self.portada()
        self.secciones()
        self.espacios_tras_bloques()
        self.limpiar()
        self.ajustes()
        return self.doc

    def guardar(self, destino):
        self.doc.save(str(destino))

    # ------------------------------------------------------------------
    def marcadores_bm(self):
        for p, m in self.marcador(r'⟦BM:([A-Za-z0-9_]+)⟧'):
            quitar_en_texto(p, r'⟦BM:[A-Za-z0-9_]+⟧')
            self.marcador_bm(p, m.group(1))

    def titulos(self):
        """Mete dentro del párrafo los marcadores que pandoc pone antes de cada título."""
        for ini in list(self.body.findall(qn('w:bookmarkStart'))):
            sig = ini.getnext()
            while sig is not None and sig.tag == qn('w:bookmarkStart'):
                sig = sig.getnext()
            if sig is None or sig.tag != qn('w:p'):
                continue
            bid = ini.get(qn('w:id'))
            fin = None
            for f in self.body.iter(qn('w:bookmarkEnd')):
                if f.get(qn('w:id')) == bid:
                    fin = f
                    break
            self.eliminar(ini)
            pp = sig.find(qn('w:pPr'))
            if pp is not None:
                pp.addnext(ini)
            else:
                sig.insert(0, ini)
            if fin is not None:
                self.eliminar(fin)
                sig.append(fin)

    # ------------------------------------------------------------------
    def _borde(self, tcpr, lado, sz, color='000000', val='single'):
        bordes = tcpr.find(qn('w:tcBorders'))
        if bordes is None:
            bordes = OxmlElement('w:tcBorders')
            # tcBorders va después de tcW/gridSpan/vMerge
            despues = None
            for t in ('w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge'):
                h = tcpr.find(qn(t))
                if h is not None:
                    despues = h
            if despues is not None:
                despues.addnext(bordes)
            else:
                tcpr.insert(0, bordes)
        viejo = bordes.find(qn('w:' + lado))
        if viejo is not None:
            bordes.remove(viejo)
        nuevo = el('w:' + lado, w_val=val, w_sz=str(sz), w_space='0', w_color=color)
        orden = ['top', 'start', 'left', 'bottom', 'end', 'right', 'insideH', 'insideV', 'tl2br', 'tr2bl']
        for h in bordes:
            if orden.index(h.tag.split('}')[1]) > orden.index(lado):
                h.addprevious(nuevo)
                break
        else:
            bordes.append(nuevo)

    def _formato_tabla(self, tbl, info):
        anchos = info['anchos_cm']
        tw = [int(round(a * CM_TW)) for a in anchos]
        grid = tbl.find(qn('w:tblGrid'))
        cols = grid.findall(qn('w:gridCol')) if grid is not None else []
        if len(cols) != len(tw):
            self.avisos.append('Tabla con %d columnas en Word y %d en LaTeX; anchos iguales.' % (len(cols), len(tw)))
            n = max(len(cols), 1)
            tw = [ANCHO_UTIL_TW // n] * n
        for c, w in zip(cols, tw):
            c.set(qn('w:w'), str(w))
        total = sum(tw)
        tblpr = tbl.find(qn('w:tblPr'))
        for t in ('w:tblW', 'w:jc', 'w:tblLayout', 'w:tblBorders', 'w:tblCellMar', 'w:tblInd'):
            for v in tblpr.findall(qn(t)):
                tblpr.remove(v)
        orden_tblpr = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
                       'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd',
                       'tblLayout', 'tblCellMar', 'tblLook']

        def poner(e):
            idx = orden_tblpr.index(e.tag.split('}')[1])
            for h in tblpr:
                n = h.tag.split('}')[1]
                if n in orden_tblpr and orden_tblpr.index(n) > idx:
                    h.addprevious(e)
                    return
            tblpr.append(e)
        poner(el('w:tblW', w_w=str(total), w_type='dxa'))
        poner(el('w:jc', w_val='center' if total < ANCHO_UTIL_TW * 0.97 else 'left'))
        poner(el('w:tblLayout', w_type='fixed'))
        mar = OxmlElement('w:tblCellMar')
        for lado, v in (('top', 0), ('left', 85), ('bottom', 0), ('right', 85)):
            mar.append(el('w:' + lado, w_w=str(v), w_type='dxa'))
        poner(mar)
        bordes = info.get('bordes', 'apa')
        if bordes == 'cuadricula':
            b = OxmlElement('w:tblBorders')
            for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                b.append(el('w:' + lado, w_val='single', w_sz='4', w_space='0', w_color='000000'))
            poner(b)
        filas = tbl.findall(qn('w:tr'))
        n_enc = info.get('n_encabezado', 0)
        reglas = set(info.get('reglas_cuerpo', []))
        tam = self.formulario_pt if bordes == 'cuadricula' else self.tabla_pt
        for i, tr in enumerate(filas):
            trpr = tr.find(qn('w:trPr'))
            if i < n_enc:
                if trpr is None:
                    trpr = OxmlElement('w:trPr')
                    tr.insert(0, trpr)
                if trpr.find(qn('w:tblHeader')) is None:
                    trpr.append(OxmlElement('w:tblHeader'))
                if trpr.find(qn('w:cantSplit')) is None:
                    trpr.insert(0, OxmlElement('w:cantSplit'))
            col = 0
            for tc in tr.findall(qn('w:tc')):
                tcpr = tc.find(qn('w:tcPr'))
                if tcpr is None:
                    tcpr = OxmlElement('w:tcPr')
                    tc.insert(0, tcpr)
                gs = tcpr.find(qn('w:gridSpan'))
                span = int(gs.get(qn('w:val'))) if gs is not None else 1
                w = sum(tw[col:col + span]) if col < len(tw) else tw[-1]
                tcw = tcpr.find(qn('w:tcW'))
                if tcw is None:
                    tcw = OxmlElement('w:tcW')
                    tcpr.insert(0, tcw)
                tcw.set(qn('w:w'), str(w))
                tcw.set(qn('w:type'), 'dxa')
                col += span
                if bordes == 'apa':
                    if i == 0 and n_enc:
                        self._borde(tcpr, 'top', 8)
                    if n_enc and i == n_enc - 1:
                        self._borde(tcpr, 'bottom', 4)
                    if i >= n_enc and (i - n_enc) in reglas:
                        self._borde(tcpr, 'top', 4)
                    if i == len(filas) - 1 and n_enc:
                        self._borde(tcpr, 'bottom', 8)
                for p in tc.findall(qn('w:p')):
                    poner_ppr(p, 'pStyle', w_val='CeldaAPA')
                    if i < n_enc and bordes == 'apa':
                        poner_ppr(p, 'jc', w_val='center')
                    for r in p.iter(qn('w:r')):
                        rpr = r.find(qn('w:rPr'))
                        if rpr is None:
                            rpr = OxmlElement('w:rPr')
                            r.insert(0, rpr)
                        if tam != 12:
                            for t in ('w:sz', 'w:szCs'):
                                v = rpr.find(qn(t))
                                if v is None:
                                    v = OxmlElement(t)
                                    rpr.append(v)
                                v.set(qn('w:val'), str(int(tam * 2)))

    def tablas(self):
        datos = self.datos.get('tablas', [])
        for p, m in self.marcador(r'^\s*⟦TBL:(\d+)⟧\s*$'):
            k = int(m.group(1))
            sig = p.getnext()
            while sig is not None and sig.tag not in (qn('w:tbl'), qn('w:p')):
                sig = sig.getnext()
            if sig is not None and sig.tag == qn('w:tbl') and k < len(datos):
                self._formato_tabla(sig, datos[k])
            else:
                self.avisos.append('No se encontró la tabla del marcador %d en el .docx.' % k)
            self.eliminar(p)

    # ------------------------------------------------------------------
    def _tabla_ecuacion(self):
        tbl = OxmlElement('w:tbl')
        tblpr = OxmlElement('w:tblPr')
        tblpr.append(el('w:tblW', w_w=str(ANCHO_UTIL_TW), w_type='dxa'))
        tblpr.append(el('w:tblLayout', w_type='fixed'))
        mar = OxmlElement('w:tblCellMar')
        for lado in ('left', 'right'):
            mar.append(el('w:' + lado, w_w='0', w_type='dxa'))
        tblpr.append(mar)
        tblpr.append(el('w:tblLook', w_val='0000', w_firstRow='0', w_lastRow='0', w_firstColumn='0',
                        w_lastColumn='0', w_noHBand='1', w_noVBand='1'))
        tblpr.append(el('w:tblCaption', w_val='Ecuación'))
        tbl.append(tblpr)
        grid = OxmlElement('w:tblGrid')
        anchos = (1134, ANCHO_UTIL_TW - 2 * 1134, 1134)
        for w in anchos:
            grid.append(el('w:gridCol', w_w=str(w)))
        tbl.append(grid)
        return tbl, anchos

    def _fila_ecuacion(self, tbl, anchos, p_math, numero):
        tr = OxmlElement('w:tr')
        trpr = OxmlElement('w:trPr')
        trpr.append(OxmlElement('w:cantSplit'))
        tr.append(trpr)
        for j, w in enumerate(anchos):
            tc = OxmlElement('w:tc')
            tcpr = OxmlElement('w:tcPr')
            tcpr.append(el('w:tcW', w_w=str(w), w_type='dxa'))
            tcpr.append(el('w:vAlign', w_val='center'))
            tc.append(tcpr)
            if j == 1:
                self.eliminar(p_math)
                pp = pPr(p_math)
                for h in list(pp):
                    pp.remove(h)
                poner_ppr(p_math, 'pStyle', w_val='EcuacionAPA')
                tc.append(p_math)
            else:
                p = OxmlElement('w:p')
                poner_ppr(p, 'pStyle', w_val='EcuacionAPA')
                if j == 2:
                    poner_ppr(p, 'jc', w_val='right')
                    r = OxmlElement('w:r')
                    t = OxmlElement('w:t')
                    t.text = '(%s)' % numero
                    r.append(t)
                    p.append(r)
                tc.append(p)
            tr.append(tc)
        tbl.append(tr)

    def ecuaciones(self):
        for p, m in self.marcador(r'^\s*⟦EQ:([^⟧]+)⟧\s*$'):
            numero = m.group(1)
            sig = p.getnext()
            if sig is None or sig.tag != qn('w:p') or not list(sig.iter(qn('m:oMathPara'))):
                self.avisos.append('No se encontró la ecuación (%s) tras su marcador.' % numero)
                self.eliminar(p)
                continue
            previo = p.getprevious()
            if previo is not None and es_tabla_ecuacion(previo):
                tbl = previo
                anchos = (1134, ANCHO_UTIL_TW - 2 * 1134, 1134)
                self._fila_ecuacion(tbl, anchos, sig, numero)
                self.eliminar(p)
                continue
            tbl, anchos = self._tabla_ecuacion()
            p.addprevious(tbl)
            self._fila_ecuacion(tbl, anchos, sig, numero)
            self.eliminar(p)
        # ecuaciones sin número: párrafo centrado sin sangría
        for p in self.body.iter(qn('w:p')):
            if list(p.iter(qn('m:oMathPara'))) and estilo_de(p) in ('BodyText', 'FirstParagraph', None):
                poner_ppr(p, 'pStyle', w_val='EcuacionAPA')

    # ------------------------------------------------------------------
    def sin_sangria(self):
        for p, m in self.marcador(r'⟦NI⟧'):
            quitar_en_texto(p, r'⟦NI⟧\s*')
            # quitar espacio inicial que deja el marcador
            for t in p.iter(qn('w:t')):
                if t.text:
                    t.text = t.text.lstrip()
                    break
            if not texto_p(p).strip() and not list(p.iter(qn('m:oMath'))):
                self.eliminar(p)
                continue
            if estilo_de(p) in ('BodyText', 'FirstParagraph', None):
                poner_ppr(p, 'ind', w_firstLine='0')

    def tabuladores(self):
        for p, m in self.marcador(r'⟦TAB⟧'):
            for t in list(p.iter(qn('w:t'))):
                if t.text and '⟦TAB⟧' in t.text:
                    r = t.getparent()
                    partes = t.text.split('⟦TAB⟧')
                    t.text = partes[0]
                    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    ancla = t
                    for parte in partes[1:]:
                        tab = OxmlElement('w:tab')
                        ancla.addnext(tab)
                        nt = OxmlElement('w:t')
                        nt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        nt.text = parte.lstrip(' ')
                        tab.addnext(nt)
                        ancla = nt

    # ------------------------------------------------------------------
    def _campo_pageref(self, destino_bm):
        runs = []
        for tipo in ('begin', 'instr', 'separate', 'valor', 'end'):
            r = OxmlElement('w:r')
            if tipo == 'instr':
                it = OxmlElement('w:instrText')
                it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                it.text = ' PAGEREF %s \\h ' % destino_bm
                r.append(it)
            elif tipo == 'valor':
                t = OxmlElement('w:t')
                t.text = self.datos.get('paginas', {}).get(destino_bm, '')
                r.append(t)
            else:
                r.append(el('w:fldChar', w_fldCharType=tipo))
            runs.append(r)
        return runs

    def indices(self):
        for p, m in self.marcador(r'⟦TOCREF:([A-Za-z0-9_]+)⟧'):
            destino = m.group(1)
            quitar_en_texto(p, r'⟦TOCREF:[A-Za-z0-9_]+⟧')
            hl = el('w:hyperlink', w_anchor=destino, w_history='1')
            for r in [h for h in p if h.tag in (qn('w:r'), qn('w:hyperlink'))]:
                self.eliminar(r)
                hl.append(r)
            tab = OxmlElement('w:r')
            tab.append(OxmlElement('w:tab'))
            hl.append(tab)
            for r in self._campo_pageref(destino):
                hl.append(r)
            p.append(hl)

    # ------------------------------------------------------------------
    def imagenes(self):
        for p in self.body.iter(qn('w:p')):
            if list(p.iter(qn('w:drawing'))) and estilo_de(p) in ('BodyText', 'FirstParagraph', 'Compact', None,
                                                                   'CaptionedFigure', 'Figure'):
                poner_ppr(p, 'pStyle', w_val='FiguraAPA')

    # ------------------------------------------------------------------
    def portada(self):
        fin = self.marcador(r'⟦FINCARATULA⟧')
        if not fin:
            return
        p_fin = fin[0][0]
        bloque = []
        e = self.body[0]
        while e is not None and e is not p_fin:
            bloque.append(e)
            e = e.getnext()
        cm_pt = 72 / 2.54
        alto_linea = 0.97      # cm, interlineado doble con Times 12
        total = 0.0
        vfills = []
        pendiente_espacio = None
        for e in bloque:
            if e.tag != qn('w:p'):
                continue
            t = texto_p(e)
            mv = re.match(r'^\s*⟦VSPACE:([0-9.]+)⟧\s*$', t)
            mf = re.match(r'^\s*⟦VFILL⟧\s*$', t)
            if mv:
                pendiente_espacio = ('fijo', float(mv.group(1)))
                total += float(mv.group(1))
                self.eliminar(e)
                continue
            if mf:
                pendiente_espacio = ('vfill', None)
                self.eliminar(e)
                continue
            dibujos = list(e.iter(qn('wp:extent')))
            if dibujos:
                total += sum(int(d.get('cy')) / 360000 for d in dibujos) + 0.4
            else:
                negrita = bool(list(e.iter(qn('w:b'))))
                por_linea = 78 if negrita else 84
                lineas = max(1, math.ceil(len(t) / por_linea))
                total += lineas * alto_linea
            if pendiente_espacio:
                if pendiente_espacio[0] == 'fijo':
                    poner_ppr(e, 'spacing', w_before=str(int(pendiente_espacio[1] * cm_pt * 20)))
                else:
                    vfills.append(e)
                pendiente_espacio = None
        disponible = 22.86 - 1.2 - total
        if vfills and disponible > 0:
            cada = max(0.4, disponible / len(vfills))
            for e in vfills:
                poner_ppr(e, 'spacing', w_before=str(int(cada * cm_pt * 20)))
        self.eliminar(p_fin)

    # ------------------------------------------------------------------
    def _sectpr_base(self):
        sp = self.body.find(qn('w:sectPr'))
        return sp

    def _nuevo_sectpr(self, cabecera=True, numeracion=None, inicio=None, horizontal=False):
        base = copy.deepcopy(self._sectpr_base())
        for t in ('w:headerReference', 'w:footerReference', 'w:pgNumType', 'w:titlePg', 'w:type'):
            for h in base.findall(qn(t)):
                base.remove(h)
        if cabecera:
            for h in reversed(self._sectpr_base().findall(qn('w:headerReference'))):
                base.insert(0, copy.deepcopy(h))
        pgsz = base.find(qn('w:pgSz'))
        if horizontal and pgsz is not None:
            pgsz.set(qn('w:w'), '15840')
            pgsz.set(qn('w:h'), '12240')
            pgsz.set(qn('w:orient'), 'landscape')
        if numeracion:
            pn = el('w:pgNumType', w_fmt=numeracion)
            if inicio is not None:
                pn.set(qn('w:start'), str(inicio))
            pgmar = base.find(qn('w:pgMar'))
            ref = base.find(qn('w:cols'))
            if ref is not None:
                ref.addprevious(pn)
            elif pgmar is not None:
                pgmar.addnext(pn)
            else:
                base.append(pn)
        return base

    def _cerrar_seccion(self, marca, sectpr):
        """Pone el sectPr en el último párrafo antes de 'marca' (o en la marca misma)."""
        previo = marca.getprevious()
        while previo is not None and previo.tag not in (qn('w:p'), qn('w:tbl')):
            previo = previo.getprevious()
        if previo is not None and previo.tag == qn('w:p') and pPr(previo).find(qn('w:sectPr')) is None:
            pPr(previo).append(sectpr)
            self.eliminar(marca)
            return previo
        for r in list(marca):
            if r.tag != qn('w:pPr'):
                marca.remove(r)
        poner_ppr(marca, 'spacing', w_before='0', w_after='0', w_line='20', w_lineRule='exact')
        pp = pPr(marca)
        pp.append(sectpr)
        return marca

    def _sin_salto_inicial(self, p_fin_seccion):
        sig = p_fin_seccion.getnext()
        while sig is not None and sig.tag != qn('w:p'):
            sig = sig.getnext()
        if sig is not None:
            poner_ppr(sig, 'pageBreakBefore', w_val='0')

    def secciones(self):
        base = self._sectpr_base()
        if base is None:
            self.avisos.append('El documento no tiene sectPr; no se crean secciones.')
            return
        # páginas de PDF: verticales con salto de página; horizontales en secciones apaisadas
        grupos = []
        actual = None
        for p, m in self.marcador(r'^\s*⟦PDFPAG:([HV])⟧\s*$'):
            img = p.getnext()
            while img is not None and img.tag != qn('w:p'):
                img = img.getnext()
            if m.group(1) == 'V':
                if img is not None:
                    poner_ppr(img, 'pageBreakBefore')
                self.eliminar(p)
                actual = None
                continue
            if actual is not None and actual['ultima'] is not None:
                # ¿consecutiva a la anterior página horizontal?
                e = actual['ultima'].getnext()
                while e is not None and e is not p and e.tag == qn('w:p') and not texto_p(e).strip() \
                        and not list(e.iter(qn('w:drawing'))):
                    e = e.getnext()
                if e is p:
                    actual['imagenes'].append(img)
                    actual['ultima'] = img
                    self.eliminar(p)
                    continue
            actual = {'marca': p, 'imagenes': [img], 'ultima': img}
            grupos.append(actual)

        num_romano = self.marcador(r'^\s*⟦SECCION:roman⟧\s*$')
        num_arabigo = self.marcador(r'^\s*⟦SECCION:arabic⟧\s*$')
        if num_romano:
            fin = self._cerrar_seccion(num_romano[0][0], self._nuevo_sectpr(cabecera=False))
            self._sin_salto_inicial(fin)
        if num_arabigo:
            fin = self._cerrar_seccion(num_arabigo[0][0], self._nuevo_sectpr(numeracion='lowerRoman', inicio=1))
            self._sin_salto_inicial(fin)
        primera_arabiga = True
        for g in grupos:
            fin = self._cerrar_seccion(g['marca'], self._nuevo_sectpr(
                numeracion='decimal', inicio=1 if primera_arabiga else None))
            primera_arabiga = False
            for j, img in enumerate(g['imagenes']):
                if j > 0:
                    poner_ppr(img, 'pageBreakBefore')
                else:
                    poner_ppr(img, 'pageBreakBefore', w_val='0')
            ultima = g['imagenes'][-1]
            pPr(ultima).append(self._nuevo_sectpr(numeracion='decimal', horizontal=True))
            self._sin_salto_inicial(ultima)
        # sección final (cuerpo): numeración arábiga
        for t in ('w:pgNumType',):
            for h in base.findall(qn(t)):
                base.remove(h)
        pn = el('w:pgNumType', w_fmt='decimal')
        if primera_arabiga and num_arabigo:
            pn.set(qn('w:start'), '1')
        ref = base.find(qn('w:cols'))
        if ref is not None:
            ref.addprevious(pn)
        else:
            base.append(pn)

    # ------------------------------------------------------------------
    def espacios_tras_bloques(self):
        """Línea en blanco (12 pt) entre una tabla/figura y el texto que sigue (APA, p. 19)."""
        for e in list(self.body):
            es_bloque = (e.tag == qn('w:tbl') and not es_tabla_ecuacion(e)) or \
                        (e.tag == qn('w:p') and estilo_de(e) == 'FiguraAPA')
            if not es_bloque:
                continue
            sig = e.getnext()
            if sig is None or sig.tag != qn('w:p'):
                continue
            if estilo_de(sig) in ('NotaAPA', 'FiguraAPA', 'Caption', 'Heading1', 'Heading2', 'Heading3',
                                  'TOCHeading'):
                continue
            sp = pPr(sig).find(qn('w:spacing'))
            if sp is None or not sp.get(qn('w:before')):
                poner_ppr(sig, 'spacing', w_before='240')

    def limpiar(self):
        for p, m in self.marcador(r'^\s*⟦[A-Z]+(:[^⟧]*)?⟧\s*$'):
            self.eliminar(p)
        restantes = self.marcador(r'⟦[A-Z]+')
        for p, m in restantes:
            self.avisos.append('Marcador sin procesar en el .docx: «%s»' % texto_p(p)[:80])
            quitar_en_texto(p, r'⟦[A-Z]+(:[^⟧]*)?⟧')
        # párrafos vacíos sin contenido dentro de la carátula y antes de títulos
        for p in list(self.body.iter(qn('w:p'))):
            if p.getparent() is not self.body:
                continue
            if texto_p(p).strip() or list(p.iter(qn('w:drawing'))) or list(p.iter(qn('m:oMath'))) \
                    or p.find(qn('w:pPr')) is not None and pPr(p).find(qn('w:sectPr')) is not None \
                    or list(p.iter(qn('w:bookmarkStart'))) or list(p.iter(qn('w:fldChar'))):
                continue
            if estilo_de(p) in ('BodyText', 'FirstParagraph', 'Compact', None, 'Caratula', 'NotaAPA'):
                self.eliminar(p)

    def ajustes(self):
        ajustes = self.doc.settings.element
        if ajustes.find(qn('w:updateFields')) is None:
            ajustes.append(el('w:updateFields', w_val='true'))
        props = self.doc.core_properties
        meta = self.datos.get('metadatos', {})
        props.title = meta.get('titulo', '')
        props.subject = meta.get('asunto', '')
        props.author = meta.get('autor', '')
        props.language = 'es-BO'
        props.keywords = meta.get('palabras_clave', '')
        props.comments = 'Generado con 07_Entregables/scripts/tesis_a_docx.py a partir de la fuente LaTeX.'
